"""Акт сверки в раскладке 1С для тестов — без настоящих данных клиентов.

Повторяет выгрузку «Акт сверки № … от …» из 1С: две половины «По данным …»,
документ в объединённых C:D, суммы числами, дата строкой «10.06.26», строка
счёта-фактуры без сумм после каждой реализации.
"""

from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook

SALE = "Реализация ТМЗ и услуг {number} от {day}"
BANK = "Платежное поручение (входящее) {number} от {day}, Накладная/Чек № вх. {number} от {day}"
CASH = "Приходный кассовый ордер {number} от {day}"
INVOICE = "Счет-фактура (выданный) {number} от {day}"


def sale(day: str, amount, number: str = "00000000527") -> tuple:
    """Реализация ``day`` в виде «10.06.2026» на ``amount`` (дебет)."""
    return (day, SALE.format(number=number, day=day), amount, None)


def bank(day: str, amount, number: str = "00000000394") -> tuple:
    """Входящее платёжное поручение (кредит)."""
    return (day, BANK.format(number=number, day=day), None, amount)


def cash(day: str, amount, number: str = "00000000012") -> tuple:
    """Приходный кассовый ордер (кредит)."""
    return (day, CASH.format(number=number, day=day), None, amount)


def _short(day: str) -> str:
    """«10.06.2026» → «10.06.26», как 1С пишет дату строки."""
    return f"{day[:6]}{day[-2:]}"


def _total(values) -> Decimal:
    """Сумма ячеек: числа или текст «6 625,00», как 1С пишет его в ячейку."""
    return sum(
        (Decimal("".join(str(value).split()).replace(",", ".")) for value in values if value is not None),
        Decimal("0"),
    )


def act_xlsx(
    rows,
    *,
    currency: str = "USD",
    counterparty: str = 'ОсОО "Тест Агро"',
    number: str = "00000000509",
    period=("01.01.2026", "07.10.2026"),
    opening=(None, None),
    turnover=None,
    closing=None,
    header=("Дата", "Документ", "Дебет", "Кредит"),
    with_invoices: bool = True,
) -> bytes:
    """Файл акта: ``rows`` — кортежи (день, документ, дебет, кредит) из :func:`sale` и др.

    ``turnover``/``closing`` по умолчанию — честные итоги строк; тест подменяет
    их, чтобы проверить отказ на расхождении.
    """
    debit = _total(row[2] for row in rows)
    credit = _total(row[3] for row in rows)
    if turnover is None:
        turnover = (debit, credit)
    if closing is None:
        balance = _total([opening[0]]) - _total([opening[1]]) + debit - credit
        closing = (balance, None) if balance > 0 else (None, -balance if balance else None)

    wb = Workbook()
    ws = wb.active
    ws.title = "Лист_1"
    ours = 'Товарищество с ограниченной ответственностью "Асыл LTD"'
    ws["B2"] = f"Акт сверки № {number} от 7 октября 2026 г."
    ws["B3"] = (
        f"взаимных расчетов за период с {period[0]} по {period[1]} между {ours} и {counterparty} "
        "по договору Контракт №05 от 08.06.2026 года"
    )
    ws["B7"] = f"По данным {ours}, {currency}"
    ws["I7"] = f"По данным {counterparty}, {currency}"
    for column, title in zip(("B", "C", "E", "F"), header):
        ws[f"{column}8"] = title
    for column, title in zip(("I", "K", "L", "N"), ("Дата", "Документ", "Дебет", "Кредит")):
        ws[f"{column}8"] = title
    ws["B9"] = "Сальдо на начало 1 января 2026 г."
    ws["I9"] = "Сальдо на начало 1 января 2026 г."
    ws["E9"], ws["F9"] = opening

    line = 10
    for invoice_number, (day, document, row_debit, row_credit) in enumerate(rows, start=552):
        ws[f"B{line}"] = _short(day) if isinstance(day, str) else day
        ws[f"C{line}"] = document
        ws.merge_cells(f"C{line}:D{line}")
        ws[f"E{line}"] = row_debit
        ws[f"F{line}"] = row_credit
        line += 1
        if with_invoices and row_debit is not None and isinstance(day, str):
            ws[f"C{line}"] = INVOICE.format(number=invoice_number, day=day)
            ws.merge_cells(f"C{line}:D{line}")
            line += 1
    ws[f"B{line}"] = "Обороты за период"
    ws[f"I{line}"] = "Обороты за период"
    ws[f"E{line}"], ws[f"F{line}"] = turnover
    line += 1
    ws[f"B{line}"] = "Сальдо на конец 7 октября 2026 г."
    ws[f"I{line}"] = "Сальдо на конец 7 октября 2026 г."
    ws[f"E{line}"], ws[f"F{line}"] = closing
    ws[f"B{line + 3}"] = f"на 07.10.2026 задолженность в пользу {counterparty}"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
