"""Разбор xlsx «Акт сверки взаимных расчётов», выгруженного из 1С.

Акт — две половины: «По данным <мы>, USD» слева и «По данным <контрагент>,
USD» справа. Читаем только свою (левую): строки «Дата | Документ | Дебет |
Кредит» между «Сальдо на начало» и «Обороты за период». Дебет — реализация
(клиент должен нам), кредит — поступление денег. Строки без сумм (счета-фактуры)
пропускаются.

Разбор строгий: колонки ищутся по заголовкам, суммы строк обязаны сойтись с
«Обороты за период», а сальдо начала с оборотами — с «Сальдо на конец». Любое
расхождение или документ незнакомого вида — отказ со списком строк: молча
записать возврат товара оплатой хуже, чем не загрузить акт.
"""

import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from rest_framework.exceptions import ValidationError

from apps.common.money import ZERO

# Выгрузка 1С за год — сотни строк; больше — уже не акт одного клиента.
MAX_ROWS = 20_000
# Сколько проблемных строк перечислить в отказе — остальные «и ещё N».
_LISTED_PROBLEMS = 10

_CURRENCIES = {"USD": "USD", "KZT": "KZT", "ТЕНГЕ": "KZT", "ТГ": "KZT"}
_SIDE = re.compile(r"^По данным (?P<name>.+),\s*(?P<currency>[^,]+?)\s*$")
_TITLE = re.compile(r"Акт сверки\s*№\s*(?P<number>[^\s,]+)")
_PERIOD = re.compile(r"за период с (\d{2}\.\d{2}\.\d{4}) по (\d{2}\.\d{2}\.\d{4})")
_COLUMNS = ("Дата", "Документ", "Дебет", "Кредит")
# Документ без хвоста «, Накладная/Чек № вх. …»: до первой даты «от ДД.ММ.ГГГГ».
_DOCUMENT_KEY = re.compile(r"^.*?\bот \d{2}\.\d{2}\.\d{4}")

# Вид документа по началу текста: (шаблон, сторона суммы, вид, способ оплаты).
_DOCUMENT_KINDS = (
    (re.compile(r"^Реализация\b"), "debit", "sale", None),
    (re.compile(r"^Платежное поручение \(?входящее\)?"), "credit", "payment", "bank"),
    (re.compile(r"^Поступление (на расчетный счет|безналичных)"), "credit", "payment", "bank"),
    (re.compile(r"^Приходный кассовый ордер\b"), "credit", "payment", "cash"),
)


@dataclass(frozen=True)
class ActRow:
    """Документ акта с суммой: реализация (``sale``) или поступление денег (``payment``)."""

    line: int
    day: date
    kind: str
    document: str
    amount: Decimal
    # Способ оплаты CRM для поступления: ``bank`` или ``cash``.
    method: str | None


@dataclass(frozen=True)
class ParsedAct:
    number: str
    period_start: date
    period_end: date
    counterparty: str
    currency: str
    # Сальдо со знаком: больше нуля — клиент должен нам, меньше — мы клиенту.
    opening: Decimal
    closing: Decimal
    rows: tuple[ActRow, ...]


def _refuse(detail: str):
    raise ValidationError({"detail": detail, "code": "act_format"})


def _text(value) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def _plain(text: str) -> str:
    """Для сравнения видов документов: «ё» 1С пишет то так, то эдак."""
    return text.replace("ё", "е").replace("Ё", "Е")


def _amount(value) -> Decimal:
    """Сумма ячейки: число или текст «6 625,00»; пусто — ноль. Ошибка — ``ValueError``."""
    if value is None or value == "":
        return ZERO
    if isinstance(value, bool):
        raise ValueError(value)
    if isinstance(value, (int, float, Decimal)):
        raw = str(value)
    else:
        raw = "".join(str(value).split()).replace(",", ".")
    try:
        amount = Decimal(raw)
    except InvalidOperation as exc:
        raise ValueError(value) from exc
    if not amount.is_finite():
        raise ValueError(value)
    return amount.quantize(Decimal("0.01"))


def _day(value) -> date:
    """День строки: ячейка-дата или текст «10.06.26» / «10.06.2026». Ошибка — ``ValueError``."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    for pattern in ("%d.%m.%y", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    raise ValueError(value)


def _load_rows(content: bytes) -> list[tuple]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(BytesIO(content), data_only=True)
    except Exception:  # noqa: BLE001 — openpyxl бросает что угодно на чужом файле
        _refuse("Файл не открывается как Excel (.xlsx) — выгрузите акт из 1С в формате xlsx")
    sheet = workbook.worksheets[0]
    if sheet.max_row > MAX_ROWS:
        _refuse(f"В файле больше {MAX_ROWS} строк — это не акт сверки одного клиента")
    return list(sheet.iter_rows(values_only=True))


def _find(rows, pattern) -> re.Match | None:
    for row in rows:
        for value in row:
            if isinstance(value, str) and (match := pattern.search(_text(value))):
                return match
    return None


def _sides(rows) -> tuple[int, dict, tuple[int, int]]:
    """Строка «По данным …»: её индекс, своя половина и границы её колонок."""
    for index, row in enumerate(rows):
        sides = [
            (column, match)
            for column, value in enumerate(row)
            if isinstance(value, str) and (match := _SIDE.match(_text(value)))
        ]
        if sides:
            right = sides[1][0] if len(sides) > 1 else len(row)
            counterparty = sides[1][1]["name"] if len(sides) > 1 else ""
            return index, {"match": sides[0][1], "counterparty": counterparty}, (sides[0][0], right)
    _refuse("Не найдена шапка «По данным …, валюта» — это не акт сверки из 1С")


def _currency(label: str) -> str:
    currency = _CURRENCIES.get(label.strip().upper())
    if currency is None:
        _refuse(f"Неизвестная валюта акта «{label}» — поддерживаются тенге (KZT) и доллары (USD)")
    return currency


def _columns(rows, start: int, bounds: tuple[int, int]) -> tuple[int, dict[str, int]]:
    """Строка заголовков своей половины и колонка каждого из «Дата/Документ/Дебет/Кредит»."""
    low, high = bounds
    for index in range(start, min(start + 5, len(rows))):
        found = {}
        for column, value in enumerate(rows[index][low:high], start=low):
            title = _text(value)
            if title in _COLUMNS and title not in found:
                found[title] = column
        if "Дата" in found or "Документ" in found:
            missing = [title for title in _COLUMNS if title not in found]
            if missing:
                _refuse(f"В акте нет колонок {', '.join(f'«{t}»' for t in missing)}: нужны {', '.join(_COLUMNS)}")
            return index, found
    _refuse("Не найдена строка заголовков «Дата | Документ | Дебет | Кредит»")


def _classify(document: str, debit: Decimal, credit: Decimal) -> tuple[str, str | None] | None:
    side = "debit" if debit and not credit else "credit" if credit and not debit else None
    plain = _plain(document)
    for pattern, kind_side, kind, method in _DOCUMENT_KINDS:
        if pattern.match(plain):
            return (kind, method) if side == kind_side else None
    return None


def _document_key(document: str) -> str:
    match = _DOCUMENT_KEY.match(document)
    return (match.group(0) if match else document)[:255]


def _problems_text(problems: list[str]) -> str:
    listed = problems[:_LISTED_PROBLEMS]
    tail = len(problems) - len(listed)
    return "; ".join(listed) + (f"; и ещё {tail}" if tail else "")


def parse_act(content: bytes) -> ParsedAct:
    """Акт сверки из xlsx-файла 1С; любое нарушение раскладки — ``ValidationError(code=act_format)``."""
    rows = _load_rows(content)
    side_index, side, bounds = _sides(rows)
    currency = _currency(side["match"]["currency"])
    period = _find(rows, _PERIOD)
    if period is None:
        _refuse("Не найден период акта («за период с ДД.ММ.ГГГГ по ДД.ММ.ГГГГ»)")
    title = _find(rows, _TITLE)
    header_index, columns = _columns(rows, side_index + 1, bounds)

    def cell(row, title):
        column = columns[title]
        return row[column] if column < len(row) else None

    opening = closing = turnover = None
    movements: list[ActRow] = []
    problems: list[str] = []
    # Документ переносится один раз (реестр — по его тексту): дубль в самом
    # акте не угадать, отказываем.
    lines_by_document: dict[str, int] = {}
    for index in range(header_index + 1, len(rows)):
        row, line = rows[index], index + 1
        label = _text(cell(row, "Дата"))
        try:
            debit, credit = _amount(cell(row, "Дебет")), _amount(cell(row, "Кредит"))
        except ValueError:
            problems.append(f"строка {line}: сумма не число")
            continue
        if label.startswith("Сальдо на начало"):
            opening = debit - credit
        elif label.startswith("Обороты за период"):
            turnover = (debit, credit)
        elif label.startswith("Сальдо на конец"):
            closing = debit - credit
            break
        elif debit or credit:
            document = _text(cell(row, "Документ"))
            try:
                day = _day(cell(row, "Дата"))
            except ValueError:
                problems.append(f"строка {line}: дата «{label}» не читается")
                continue
            kind = _classify(document, debit, credit)
            if kind is None or debit < 0 or credit < 0:
                problems.append(f"строка {line}: «{document or 'без названия'}» — такой документ не переносится")
                continue
            key = _document_key(document)
            if key in lines_by_document:
                problems.append(f"строка {line}: «{key}» уже есть в строке {lines_by_document[key]}")
                continue
            lines_by_document[key] = line
            movements.append(ActRow(
                line=line, day=day, kind=kind[0], document=key, amount=debit or credit, method=kind[1],
            ))
    if problems:
        _refuse(f"Акт не загружен: {_problems_text(problems)}")
    if turnover is None or closing is None:
        _refuse("В акте нет строк «Обороты за период» и «Сальдо на конец» — файл обрезан?")
    if not movements:
        _refuse("В акте нет документов с суммами — переносить нечего")

    debit_total = sum((row.amount for row in movements if row.kind == "sale"), ZERO)
    credit_total = sum((row.amount for row in movements if row.kind == "payment"), ZERO)
    if (debit_total, credit_total) != turnover:
        _refuse(
            f"Суммы строк ({debit_total} / {credit_total}) не сходятся с «Обороты за период» "
            f"({turnover[0]} / {turnover[1]}) — акт выгружен не полностью?"
        )
    opening = opening or ZERO
    if opening + debit_total - credit_total != closing:
        _refuse(
            f"Сальдо на начало {opening} и обороты не дают «Сальдо на конец» {closing} — проверьте акт"
        )
    return ParsedAct(
        number=title["number"] if title else "",
        period_start=datetime.strptime(period.group(1), "%d.%m.%Y").date(),
        period_end=datetime.strptime(period.group(2), "%d.%m.%Y").date(),
        counterparty=side["counterparty"],
        currency=currency,
        opening=opening,
        closing=closing,
        rows=tuple(movements),
    )
