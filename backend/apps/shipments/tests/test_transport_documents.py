"""Тягач и прицеп в документах: накладная и выписка показывают номер для людей."""
from io import BytesIO

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from apps.catalog.models import Product
from apps.clients.models import Client
from apps.orders.models import Order, OrderItem
from apps.shipments import waybill
from apps.shipments.models import Shipment, ShipmentWagon

pytestmark = pytest.mark.django_db


def _order(**fields):
    client = Client.objects.create_with_user(first_name="Азамат", last_name="К", phone="docs")
    product = Product.objects.create(name="Мука", color="Red", weight_kg="50")
    order = Order.objects.create(client=client, status="shipped", **fields)
    OrderItem.objects.create(order=order, product=product, quantity=2, unit_price="100")
    return order


def _waybill_texts(order, monkeypatch):
    texts = []
    original = waybill.Paragraph

    def capture(text, *args, **kwargs):
        texts.append(text)
        return original(text, *args, **kwargs)

    monkeypatch.setattr(waybill, "Paragraph", capture)
    assert waybill.build_waybill_pdf(order).startswith(b"%PDF")
    return texts


def test_waybill_shows_truck_and_trailer(monkeypatch):
    order = _order(truck_number="07KG695ADT", trailer_number="07KG837PB")

    texts = _waybill_texts(order, monkeypatch)

    assert "№ Автомашины: 07 KG 695 ADT / прицеп 07 KG 837 PB" in texts


@pytest.mark.parametrize(
    ("fields", "line"),
    [
        ({"truck_number": "403BJN13"}, "№ Автомашины: 403 BJN 13"),
        ({"truck_number": ""}, "№ Автомашины: —"),
        ({"transport_type": "train", "truck_number": "00123456"}, "№ Вагона: 00123456"),
        ({"transport_type": "train", "truck_number": ""}, "№ Вагона: вагон"),
    ],
)
def test_waybill_transport_line_without_trailer(monkeypatch, fields, line):
    assert line in _waybill_texts(_order(**fields), monkeypatch)


def test_statement_puts_trailer_into_the_number_cell(auth_client, user_with_perms):
    reporter = user_with_perms("statement-plates", codes=["clients.view", "reports.export"])
    order = _order(truck_number="07KG695ADT", trailer_number="07KG837PB")

    client_sheet = load_workbook(BytesIO(
        auth_client(reporter).get(f"/api/clients/{order.client_id}/statement/").content))["Заказы"]
    all_sheet = load_workbook(BytesIO(
        auth_client(reporter).get("/api/clients/statement/").content))["Заказы"]

    assert client_sheet["H3"].value == "Номер"
    assert client_sheet["H4"].value == "07 KG 695 ADT / 07 KG 837 PB"
    assert all_sheet["J3"].value == "Номер"
    assert all_sheet["J4"].value == "07 KG 695 ADT / 07 KG 837 PB"


def _report_order(numbers=("28087658", "28087666")):
    """Вагонный заказ, отгруженный по отчёту: станция и вагоны вместо номера."""
    order = _order(transport_type="train", rail_station="Раустан")
    shipment = Shipment.objects.create(order=order, bags_loaded=2, shipped_at=timezone.now())
    product = order.items.get().product
    for position, number in enumerate(numbers, start=1):
        ShipmentWagon.objects.create(
            shipment=shipment, number=number, product=product, bags=1, weight_kg="50", position=position)
    return order


def test_waybill_of_a_report_shipment_lists_station_and_wagons(monkeypatch):
    texts = _waybill_texts(_report_order(), monkeypatch)

    assert "Вагоны: 2 — список ниже" in texts
    assert "Станция назначения: Раустан" in texts
    assert {"28087658", "28087666", "Вагонов: 2"} <= set(texts)
    assert not any(text.startswith("№ Вагона") for text in texts)


def test_statement_number_cell_lists_the_wagons(auth_client, user_with_perms):
    reporter = user_with_perms("statement-wagons", codes=["clients.view", "reports.export"])
    order = _report_order()

    sheet = load_workbook(BytesIO(
        auth_client(reporter).get(f"/api/clients/{order.client_id}/statement/").content))["Заказы"]

    assert sheet["H4"].value == "28087658, 28087666"
